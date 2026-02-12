"""
Export and reporting module.
数据导出：支持 JSON / CSV / Excel 格式输出映射结果
"""

import csv
import json
import logging
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional

from .models import ExecutiveSchoolMapping

logger = logging.getLogger(__name__)


class ReportExporter(ABC):
    """Abstract base for report exporters."""

    @property
    @abstractmethod
    def file_extension(self) -> str:
        ...

    @abstractmethod
    def export(
        self,
        results: list[ExecutiveSchoolMapping],
        output_path: str,
    ) -> None:
        ...


class JSONExporter(ReportExporter):
    """Export mappings as JSON."""

    @property
    def file_extension(self) -> str:
        return "json"

    def export(
        self,
        results: list[ExecutiveSchoolMapping],
        output_path: str,
    ) -> None:
        data = [r.to_dict() for r in results]
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info("JSON export: %d records -> %s", len(data), output_path)


# CSV column headers for the flat export
_CSV_HEADERS = [
    "企业名称",
    "统一社会信用代码",
    "行业",
    "省份",
    "招标/中标项目",
    "招标类型",
    "金额",
    "高管姓名",
    "高管职位",
    "商学院",
    "项目类型",
    "届别",
    "置信度",
    "数据来源",
    "证据链接",
]


def _flatten_mapping(mapping: ExecutiveSchoolMapping) -> list[dict]:
    """Flatten a mapping into one row per school record (or one row if none)."""
    base = {
        "企业名称": mapping.company.name,
        "统一社会信用代码": mapping.company.unified_social_credit_code,
        "行业": mapping.company.industry,
        "省份": mapping.company.province,
        "招标/中标项目": mapping.company.bid_project_name,
        "招标类型": mapping.company.bid_type,
        "金额": mapping.company.bid_amount,
        "高管姓名": mapping.executive.name,
        "高管职位": mapping.executive.role.value,
    }
    if not mapping.school_records:
        return [{
            **base,
            "商学院": "",
            "项目类型": "",
            "届别": "",
            "置信度": "",
            "数据来源": "",
            "证据链接": "",
        }]

    rows = []
    for rec in mapping.school_records:
        rows.append({
            **base,
            "商学院": rec.school_name_raw or rec.school.value,
            "项目类型": rec.program,
            "届别": rec.class_year,
            "置信度": f"{rec.confidence:.2f}",
            "数据来源": rec.data_source.value,
            "证据链接": rec.evidence_url,
        })
    return rows


class CSVExporter(ReportExporter):
    """Export mappings as CSV (UTF-8 with BOM for Excel compatibility)."""

    @property
    def file_extension(self) -> str:
        return "csv"

    def export(
        self,
        results: list[ExecutiveSchoolMapping],
        output_path: str,
    ) -> None:
        rows = []
        for mapping in results:
            rows.extend(_flatten_mapping(mapping))

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=_CSV_HEADERS)
            writer.writeheader()
            for row in rows:
                writer.writerow(row)

        logger.info("CSV export: %d rows -> %s", len(rows), output_path)


class ExcelExporter(ReportExporter):
    """
    Export mappings as Excel (.xlsx).
    Requires openpyxl; falls back to CSV if unavailable.
    """

    @property
    def file_extension(self) -> str:
        return "xlsx"

    def export(
        self,
        results: list[ExecutiveSchoolMapping],
        output_path: str,
    ) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            logger.warning(
                "openpyxl not installed; falling back to CSV export"
            )
            csv_path = output_path.replace(".xlsx", ".csv")
            CSVExporter().export(results, csv_path)
            return

        rows = []
        for mapping in results:
            rows.extend(_flatten_mapping(mapping))

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "高管商学院匹配"

        # Header styling
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_font = Font(color="FFFFFF", bold=True, size=11)

        for col_idx, header in enumerate(_CSV_HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Data rows
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, header in enumerate(_CSV_HEADERS, 1):
                ws.cell(row=row_idx, column=col_idx, value=row_data.get(header, ""))

        # Auto-adjust column widths
        for col_idx, header in enumerate(_CSV_HEADERS, 1):
            max_len = len(header)
            for row_idx in range(2, len(rows) + 2):
                val = str(ws.cell(row=row_idx, column=col_idx).value or "")
                max_len = max(max_len, len(val))
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(col_idx)
            ].width = min(max_len + 4, 40)

        # Freeze header row
        ws.freeze_panes = "A2"

        # Add auto-filter
        ws.auto_filter.ref = ws.dimensions

        wb.save(output_path)
        logger.info("Excel export: %d rows -> %s", len(rows), output_path)


def generate_summary_stats(results: list[ExecutiveSchoolMapping]) -> dict:
    """
    Generate summary statistics for the matching results.
    生成匹配结果的统计摘要
    """
    total_companies = len({r.company.id_key for r in results})
    total_executives = len(results)
    matched = [r for r in results if r.school_records]
    unmatched = [r for r in results if not r.school_records]

    # School distribution
    school_count: dict[str, int] = {}
    program_count: dict[str, int] = {}
    for r in matched:
        for rec in r.school_records:
            school_label = rec.school_name_raw or rec.school.value
            school_count[school_label] = school_count.get(school_label, 0) + 1
            if rec.program:
                program_count[rec.program] = program_count.get(rec.program, 0) + 1

    # Source distribution
    source_count: dict[str, int] = {}
    for r in matched:
        for rec in r.school_records:
            src = rec.data_source.value
            source_count[src] = source_count.get(src, 0) + 1

    return {
        "total_companies": total_companies,
        "total_executives": total_executives,
        "matched_executives": len(matched),
        "unmatched_executives": len(unmatched),
        "match_rate": f"{len(matched) / max(total_executives, 1) * 100:.1f}%",
        "school_distribution": dict(
            sorted(school_count.items(), key=lambda x: -x[1])
        ),
        "program_distribution": dict(
            sorted(program_count.items(), key=lambda x: -x[1])
        ),
        "source_distribution": source_count,
    }
